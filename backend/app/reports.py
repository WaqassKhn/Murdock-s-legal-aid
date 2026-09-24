from io import BytesIO
from xml.sax.saxutils import escape

DISCLAIMER = (
    'LegalLens provides document-based legal information, not legal advice. Laws and outcomes depend on jurisdiction and circumstances. '
    'Consult a qualified legal professional before making important decisions. AI-generated results may be incomplete or incorrect.'
)


def report_sections(
    workspace: dict, documents: list[dict], objective: str, notes: str, obligations: list[dict]
):
    sections: list[tuple[str, list[str]]] = [
        (
            'Review notice',
            [
                DISCLAIMER,
                'This AI-generated informational material should be independently reviewed. Extractive findings are candidate review points, not legal conclusions.',
            ],
        ),
        ('Your objective', [objective or workspace['objective'] or 'Not provided']),
        (
            'Scope and jurisdiction',
            [
                workspace['name'],
                f'Confirmed jurisdiction: {workspace["jurisdiction"]}',
                'No external legal authorities were consulted. Missing information is not inferred.',
            ],
        ),
    ]
    for document in documents:
        analysis = document.get('analysis')
        if not analysis:
            sections.append((document['name'], ['Analysis unavailable.']))
            continue
        lines = [
            f'Source document: {document["id"]}; immutable version {document.get("version", 1)}',
            f'Analysis mode: {analysis["mode"]}',
        ]
        for key, field in analysis['metadata'].items():
            lines.append(f'{key.replace("_", " ").title()}: {field["value"]}')
            lines.extend(citation_lines(field['citations']))
        sections.append((f'Document overview — {document["name"]}', lines))
        review = []
        for risk in analysis['risks']:
            review.extend(
                [
                    f'{risk["severity"]}: {risk["finding"]}',
                    risk['why_it_matters'],
                    f'Clarify: {risk["suggested_clarification"]}',
                    f'Ask your lawyer: {risk["lawyer_question"]}',
                ]
            )
            review.extend(citation_lines(risk['citations']))
        sections.append(
            (
                f'Review points and unresolved ambiguities — {document["name"]}',
                review
                or [
                    'No candidate findings were detected. This does not establish that the document has no risks.'
                ],
            )
        )
        clauses = []
        for clause in analysis['clauses']:
            if clause['attention'] != 'Low attention':
                clauses.extend(
                    [
                        f'{clause["clause_type"]} — {clause["attention"]}',
                        f'System interpretation: {clause["explanation"]}',
                    ]
                )
                clauses.extend(citation_lines(clause['citations']))
        sections.append(
            ('High-attention and review clauses', clauses or ['No additional candidate clauses detected.'])
        )
    checklist = []
    for obligation in obligations:
        checklist.extend(
            [
                f'[{obligation["status"]}] {obligation["responsible_party"]}: {obligation["action"]}',
                f'Trigger: {obligation["trigger"]}; window: {obligation["time_window"]}; recurrence: {obligation["recurrence"]}',
                f'Consequence stated in source: {obligation["consequence"]}',
            ]
        )
        checklist.extend(citation_lines(obligation['citations']))
    sections.extend(
        [
            ('Obligations and deadlines', checklist or ['No supported obligations extracted.']),
            ('User notes (not verified evidence)', [notes or 'No notes provided.']),
        ]
    )
    return sections


def citation_lines(citations):
    return [
        f'Evidence — document {c["document_id"]}, page {c["page"]}, {c["section"]}: “{c["excerpt"]}”'
        for c in citations
    ]


def generate_report(sections, format: str) -> bytes:
    output = BytesIO()
    if format == 'docx':
        from docx import Document
        from docx.shared import Inches, Pt

        doc = Document()
        section = doc.sections[0]
        section.top_margin = section.bottom_margin = Inches(0.7)
        doc.styles['Normal'].font.name = 'Calibri'
        doc.styles['Normal'].font.size = Pt(10)
        doc.add_heading('LegalLens | Lawyer consultation pack', 0)
        for title, lines in sections:
            doc.add_heading(title, level=1)
            for line in lines:
                doc.add_paragraph(line)
        doc.save(output)
    else:
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

        styles = getSampleStyleSheet()
        styles['BodyText'].fontSize = 9
        styles['BodyText'].leading = 13
        styles['Heading1'].textColor = colors.HexColor('#163b3a')
        flow = [Paragraph('LegalLens | Lawyer consultation pack', styles['Title'])]
        for title, lines in sections:
            flow.append(Paragraph(escape(title), styles['Heading1']))
            for line in lines:
                flow.extend(
                    [Paragraph(escape(line).replace('\n', '<br/>'), styles['BodyText']), Spacer(1, 6)]
                )
        SimpleDocTemplate(output, rightMargin=44, leftMargin=44, topMargin=40, bottomMargin=40).build(flow)
    return output.getvalue()
